"""Générer le Word du TP depuis docs/Walter_Arthur_rapport.md.

Dépendance d'édition : python-docx. Aucun besoin pour exécuter le TP RSA.
Les marqueurs <!-- page --> séparent les chapitres ; vérifier le PDF après export.
"""
from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RACINE = Path(__file__).resolve().parents[1]


def texte_propre(texte):
    return texte.replace('`', '')


def tableau(document, lignes):
    cellules = [[c.strip() for c in l.strip().strip('|').split('|')] for l in lignes]
    cellules = [cellules[0]] + cellules[2:]
    n = len(cellules[0])
    table = document.add_table(rows=0, cols=n)
    table.autofit = False
    largeurs = [2, 4.9] if n == 2 else [6.9 / n] * n
    for col, w in zip(table.columns, largeurs):
        col.width = Inches(w)
    for index, valeurs in enumerate(cellules):
        row = table.add_row()
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if not index:
            pr.append(OxmlElement('w:tblHeader'))
        for cell, valeur in zip(row.cells, valeurs):
            cell.text = texte_propre(valeur)
            props = cell._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            for cote in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + cote)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    e.set(qn('w:' + key), val)
                b.append(e)
            props.append(b)
            marges = OxmlElement('w:tcMar')
            for cote in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + cote)
                m.set(qn('w:w'), '80'); m.set(qn('w:type'), 'dxa'); marges.append(m)
            props.append(marges)
            align = OxmlElement('w:vAlign'); align.set(qn('w:val'), 'center'); props.append(align)
            if not index:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EDF2'); props.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def equation(document, texte):
    par = document.add_paragraph()
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = texte
    run.append(t); math.append(run); par._p.append(math)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.8)
    for nom in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = document.styles[nom]
        style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(7)
    document.styles['Title'].font.size = Pt(25)
    document.styles['Heading 1'].font.size = Pt(19)
    document.styles['Heading 2'].font.size = Pt(13)
    footer = section.footer.paragraphs[0]; footer.alignment = 2
    r = footer.add_run('Arthur Walter · TP RSA · '); r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    document.core_properties.title = 'TP RSA Sécurité et Pratiques'
    document.core_properties.author = 'Arthur Walter'
    source = (RACINE / 'docs/Walter_Arthur_rapport.md').read_text(encoding='utf-8')
    lignes = source.splitlines(); i = 0; premier = True
    while i < len(lignes):
        ligne = lignes[i]
        if not ligne.strip():
            i += 1; continue
        if ligne == '<!-- page -->':
            document.add_page_break(); i += 1; continue
        if ligne.startswith('# '):
            document.add_heading(ligne[2:], 0 if premier else 1)
            premier = False; i += 1; continue
        if ligne.startswith('## '):
            document.add_heading(ligne[3:], 2); i += 1; continue
        if ligne.startswith('```'):
            i += 1
            while i < len(lignes) and not lignes[i].startswith('```'):
                par = document.add_paragraph()
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.line_spacing = 1
                par.paragraph_format.keep_with_next = (i + 1 < len(lignes) and not lignes[i + 1].startswith('```'))
                run = par.add_run(lignes[i]); run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(8.5)
                for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
                    run._r.get_or_add_rPr().rFonts.set(qn('w:' + key), 'DejaVu Sans Mono')
                i += 1
            document.add_paragraph().paragraph_format.space_after = Pt(0)
            i += 1; continue
        if ligne.startswith('|'):
            groupe = []
            while i < len(lignes) and lignes[i].startswith('|'):
                groupe.append(lignes[i]); i += 1
            tableau(document, groupe); continue
        if ligne.startswith('$$') and ligne.endswith('$$'):
            equation(document, ligne[2:-2]); i += 1; continue
        document.add_paragraph(texte_propre(ligne)); i += 1
    sortie = RACINE / 'docs/Walter_Arthur.docx'
    document.save(sortie)
    print(sortie)


if __name__ == '__main__':
    main()
